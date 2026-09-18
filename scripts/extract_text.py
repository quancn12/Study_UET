#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Công cụ trích xuất nội dung văn bản từ các file bài giảng/tài liệu:
Hỗ trợ: .pdf, .pptx, .docx, .txt, .md
"""

import os
import sys
from pathlib import Path

# Đảm bảo in tiếng Việt trên console Windows không bị lỗi UnicodeEncodeError
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

def extract_from_txt(file_path: Path) -> str:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()

def extract_from_docx(file_path: Path) -> str:
    try:
        import docx
        doc = docx.Document(file_path)
        content = []
        for p in doc.paragraphs:
            if p.text.strip():
                content.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    content.append(row_text)
        return "\n\n".join(content)
    except Exception as e:
        return f"[Lỗi đọc DOCX: {e}]"

def extract_from_pptx(file_path: Path) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(file_path)
        slides_text = []
        for idx, slide in enumerate(prs.slides, start=1):
            slide_content = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = "".join(run.text for run in paragraph.runs).strip()
                        if text:
                            slide_content.append(text)
            if slide_content:
                slides_text.append(f"--- [Slide {idx}] ---\n" + "\n".join(slide_content))
        return "\n\n".join(slides_text)
    except Exception as e:
        return f"[Lỗi đọc PPTX: {e}]"

def extract_from_pdf(file_path: Path) -> str:
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(file_path)
        pages_text = []
        for idx, page in enumerate(doc, start=1):
            text = page.get_text().strip()
            if text:
                pages_text.append(f"--- [Trang {idx}] ---\n{text}")
        return "\n\n".join(pages_text)
    except Exception as e:
        # Fallback to pypdf
        try:
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            pages_text = []
            for idx, page in enumerate(reader.pages, start=1):
                text = (page.extract_text() or "").strip()
                if text:
                    pages_text.append(f"--- [Trang {idx}] ---\n{text}")
            return "\n\n".join(pages_text)
        except Exception as e2:
            return f"[Lỗi đọc PDF: {e} | {e2}]"

def extract_file(file_path: Path) -> str:
    ext = file_path.suffix.lower()
    if ext in [".txt", ".md"]:
        return extract_from_txt(file_path)
    elif ext in [".docx", ".doc"]:
        return extract_from_docx(file_path)
    elif ext in [".pptx", ".ppt"]:
        return extract_from_pptx(file_path)
    elif ext == ".pdf":
        return extract_from_pdf(file_path)
    else:
        return f"[Định dạng chưa được hỗ trợ: {ext}]"

def main():
    base_dir = Path(__file__).resolve().parent.parent
    raw_dir = base_dir / "raw_materials"
    output_dir = raw_dir / "extracted"
    output_dir.mkdir(parents=True, exist_ok=True)

    if len(sys.argv) > 1:
        target_path = Path(sys.argv[1])
        files = [target_path] if target_path.is_file() else list(target_path.glob("*.*"))
    else:
        files = [f for f in raw_dir.glob("*.*") if f.is_file()]

    if not files:
        print(f"Chưa có file nào trong thư mục: {raw_dir}")
        print("Vui lòng copy file (.pdf, .pptx, .docx, .txt) vào raw_materials/ để xử lý.")
        return

    print(f"Tìm thấy {len(files)} file cần trích xuất:")
    for f in files:
        if f.suffix.lower() in [".pdf", ".pptx", ".docx", ".doc", ".txt", ".md"]:
            print(f"-> Đang xử lý: {f.name}...")
            text = extract_file(f)
            out_file = output_dir / f"{f.stem}.txt"
            with open(out_file, "w", encoding="utf-8") as out:
                out.write(text)
            print(f"   Đã lưu nội dung văn bản: {out_file.relative_to(base_dir)} ({len(text)} ký tự)")
        else:
            print(f"   Bỏ qua định dạng không hỗ trợ: {f.name}")

if __name__ == "__main__":
    main()
